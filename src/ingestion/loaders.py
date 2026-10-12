import io
from dataclasses import dataclass
from pathlib import Path
from urllib.parse import urlparse

import requests

MAX_FILE_MB = 25
SUPPORTED_EXTENSIONS = {".pdf", ".pptx", ".txt"}


class IngestionError(Exception):
    """Raised with a user-friendly message when a source can't be ingested."""


@dataclass
class Segment:
    text: str
    location: str = ""   # e.g. "Page 3", "Slide 2" (shown in citations)
    page: int = 0        # numeric page/slide, 0 if not applicable


# ---------------- PDF ----------------
def load_pdf(data: bytes) -> list[Segment]:
    from pypdf import PdfReader

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise IngestionError("This PDF is password-protected. Remove the password and try again.")
        segments = []
        for i, page in enumerate(reader.pages, start=1):
            text = (page.extract_text() or "").strip()
            if text:
                segments.append(Segment(text, f"Page {i}", i))
    except IngestionError:
        raise
    except Exception as e:
        raise IngestionError(f"Could not read this PDF: {e}")

    if not segments:
        raise IngestionError("No text found in this PDF. It may be a scanned image (OCR is not supported).")
    return segments


# ---------------- PPTX ----------------
def load_pptx(data: bytes) -> list[Segment]:
    from pptx import Presentation

    try:
        prs = Presentation(io.BytesIO(data))
        segments = []
        for i, slide in enumerate(prs.slides, start=1):
            parts = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    t = shape.text_frame.text.strip()
                    if t:
                        parts.append(t)
                if getattr(shape, "has_table", False) and shape.has_table:
                    for row in shape.table.rows:
                        parts.append(" | ".join(c.text.strip() for c in row.cells))
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
                notes = slide.notes_slide.notes_text_frame.text.strip()
                if notes:
                    parts.append(f"Speaker notes: {notes}")
            text = "\n\n".join(parts).strip()
            if text:
                segments.append(Segment(text, f"Slide {i}", i))
    except Exception as e:
        raise IngestionError(f"Could not read this PowerPoint file: {e}")

    if not segments:
        raise IngestionError("No text found in this presentation.")
    return segments


# ---------------- TXT ----------------
def load_txt(data: bytes) -> list[Segment]:
    try:
        text = data.decode("utf-8")
    except UnicodeDecodeError:
        text = data.decode("latin-1")
    text = text.strip()
    if not text:
        raise IngestionError("This text file is empty.")
    return [Segment(text)]


# ---------------- URL ----------------
def fetch_url(url: str) -> bytes:
    url = (url or "").strip()
    parsed = urlparse(url)
    if parsed.scheme not in ("http", "https") or not parsed.netloc:
        raise IngestionError("Please enter a valid URL starting with http:// or https://")
    try:
        resp = requests.get(
            url,
            timeout=15,
            headers={"User-Agent": "Mozilla/5.0 (NotebookClone/1.0)"},
        )
        resp.raise_for_status()
    except requests.RequestException as e:
        raise IngestionError(f"Could not fetch that URL: {e}")
    if len(resp.content) > MAX_FILE_MB * 1024 * 1024:
        raise IngestionError(f"That page is larger than {MAX_FILE_MB} MB.")
    return resp.content


def parse_html(html_bytes: bytes) -> tuple[list[Segment], str]:
    """Returns (segments, page_title)."""
    import trafilatura
    from bs4 import BeautifulSoup

    html = html_bytes.decode("utf-8", errors="replace")
    soup = BeautifulSoup(html, "html.parser")
    title = soup.title.get_text(strip=True) if soup.title else ""

    text = trafilatura.extract(html, include_comments=False, include_tables=True)
    if not text:  # fallback
        for tag in soup(["script", "style", "nav", "footer", "header", "aside"]):
            tag.decompose()
        text = soup.get_text("\n")
    text = "\n".join(line.strip() for line in (text or "").splitlines() if line.strip())

    if len(text) < 50:
        raise IngestionError("Could not extract readable text from that page (it may require JavaScript or a login).")
    return [Segment(text)], title


# ---------------- dispatcher ----------------
def load_file(filename: str, data: bytes) -> list[Segment]:
    ext = Path(filename).suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise IngestionError(f"Unsupported file type '{ext}'. Supported: PDF, PPTX, TXT.")
    if len(data) > MAX_FILE_MB * 1024 * 1024:
        raise IngestionError(f"File is larger than {MAX_FILE_MB} MB.")
    if ext == ".pdf":
        return load_pdf(data)
    if ext == ".pptx":
        return load_pptx(data)
    return load_txt(data)